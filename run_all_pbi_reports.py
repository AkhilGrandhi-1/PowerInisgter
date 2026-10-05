"""
Run all list_pbi_*.py extraction scripts in sequence and merge their CSV
output into a single multi-sheet Excel workbook.

Runs, in order:
  1. list_pbi_capacities.py
  2. list_pbi_workspaces.py
  3. list_pbi_datasets_and_reports.py
  4. list_pbi_tables_and_columns.py   (slowest - scans every workspace;
                                        skip with --skip-tables-columns)

Prerequisites: see PBI_ADMIN_SETUP.md and AZURE_TENANT_ID / AZURE_CLIENT_ID /
AZURE_CLIENT_SECRET in .env or as CLI args.

Credentials are passed to each sub-script via environment variables, never
as command-line arguments, so the client secret never shows up in a process
listing.

Project layout:
  run_all_pbi_reports.py   <- this file (repo root)
  requirements.txt         <- repo root
  scripts/                 <- all list_pbi_*.py + pbi_admin_common.py
  output/                  <- all generated CSVs + the combined .xlsx (default)

Usage:
  .venv/Scripts/python.exe run_all_pbi_reports.py
  .venv/Scripts/python.exe run_all_pbi_reports.py --output-dir output --skip-tables-columns
"""
import argparse
import csv
import logging
import os
import subprocess
import sys
import time
from pathlib import Path

from openpyxl import Workbook

ROOT_DIR = Path(__file__).resolve().parent
SCRIPTS_DIR = ROOT_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

from pbi_admin_common import add_credential_args, load_env, require_credentials  # noqa: E402

logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)


def run_step(name: str, script: str, script_args: list, env: dict):
    cmd = [sys.executable, str(SCRIPTS_DIR / script), *script_args]
    logger.info("=== %s ===", name)
    start = time.monotonic()
    result = subprocess.run(cmd, env=env)
    elapsed = time.monotonic() - start
    if result.returncode != 0:
        raise RuntimeError(f"{script} failed with exit code {result.returncode}")
    logger.info("%s finished in %.1fs", name, elapsed)


def csv_to_sheet(workbook: Workbook, csv_path: Path, sheet_name: str):
    if not csv_path.exists():
        logger.warning("Expected output %s not found; skipping sheet %s", csv_path, sheet_name)
        return
    sheet = workbook.create_sheet(title=sheet_name[:31])  # Excel sheet name length limit
    with open(csv_path, newline="", encoding="utf-8") as f:
        for row in csv.reader(f):
            sheet.append(row)


def main():
    load_env()

    parser = argparse.ArgumentParser(
        description="Run all Power BI admin extraction scripts and merge the results into one Excel workbook."
    )
    add_credential_args(parser)
    parser.add_argument("--output-dir", default="output", help="Directory for the CSV and Excel output.")
    parser.add_argument("--excel-output", default=None,
                         help="Path to the combined .xlsx file (default: <output-dir>/pbi_full_export.xlsx).")
    parser.add_argument("--skip-tables-columns", action="store_true",
                         help="Skip the scanner-based tables/columns extract (it's the slowest step).")
    parser.add_argument("--poll-interval", type=int, default=5,
                         help="Passed through to list_pbi_tables_and_columns.py.")
    parser.add_argument("--scan-timeout", type=int, default=600,
                         help="Passed through to list_pbi_tables_and_columns.py.")
    args = parser.parse_args()
    require_credentials(args, parser)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    excel_output = Path(args.excel_output) if args.excel_output else output_dir / "pbi_full_export.xlsx"

    child_env = dict(os.environ)
    child_env["AZURE_TENANT_ID"] = args.tenant_id
    child_env["AZURE_CLIENT_ID"] = args.client_id
    child_env["AZURE_CLIENT_SECRET"] = args.client_secret

    paths = {
        "capacities": output_dir / "pbi_capacities.csv",
        "workspaces": output_dir / "pbi_workspaces.csv",
        "datasets": output_dir / "pbi_datasets.csv",
        "reports": output_dir / "pbi_reports.csv",
        "tables": output_dir / "pbi_tables.csv",
        "columns": output_dir / "pbi_columns.csv",
    }

    overall_start = time.monotonic()

    run_step(
        "Capacities",
        "list_pbi_capacities.py",
        ["--output", str(paths["capacities"])],
        child_env,
    )
    run_step(
        "Workspaces",
        "list_pbi_workspaces.py",
        ["--output", str(paths["workspaces"])],
        child_env,
    )
    run_step(
        "Datasets & Reports",
        "list_pbi_datasets_and_reports.py",
        ["--datasets-output", str(paths["datasets"]), "--reports-output", str(paths["reports"])],
        child_env,
    )

    if args.skip_tables_columns:
        logger.info("Skipping tables/columns scan (--skip-tables-columns)")
    else:
        run_step(
            "Tables & Columns",
            "list_pbi_tables_and_columns.py",
            [
                "--tables-output", str(paths["tables"]),
                "--columns-output", str(paths["columns"]),
                "--poll-interval", str(args.poll_interval),
                "--scan-timeout", str(args.scan_timeout),
            ],
            child_env,
        )

    logger.info("Building combined workbook %s", excel_output)
    workbook = Workbook()
    workbook.remove(workbook.active)  # drop the default blank sheet

    csv_to_sheet(workbook, paths["capacities"], "Capacities")
    csv_to_sheet(workbook, paths["workspaces"], "Workspaces")
    csv_to_sheet(workbook, paths["datasets"], "Datasets")
    csv_to_sheet(workbook, paths["reports"], "Reports")
    if not args.skip_tables_columns:
        csv_to_sheet(workbook, paths["tables"], "Tables")
        csv_to_sheet(workbook, paths["columns"], "Columns")

    if not workbook.sheetnames:
        raise RuntimeError("No CSV output was produced; nothing to write to the workbook.")

    workbook.save(excel_output)

    total_elapsed = time.monotonic() - overall_start
    logger.info(
        "All done in %.1fs. Individual CSVs are in %s, combined workbook: %s",
        total_elapsed, output_dir, excel_output,
    )


if __name__ == "__main__":
    main()
